"""Turns an uploaded resume (PDF, DOCX or TXT) into text plus structured fields."""
import io
import re
from datetime import date

from .skills import extract_skills

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{2,5}\)?[\s-]?)?\d{3,5}[\s-]?\d{4,5}")
YEARS_RE = re.compile(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)(?:\s+of)?(?:\s+\w+){0,3}\s+experience", re.I)
RANGE_RE = re.compile(
    r"((?:19|20)\d{2})\s*(?:-|–|—|to)\s*((?:19|20)\d{2}|present|current|now|till date|ongoing)", re.I)
LINK_RE = re.compile(r"(?:https?://)?(?:www\.)?(linkedin\.com/in/[\w-]+|github\.com/[\w-]+)", re.I)

DEGREES = {
    "PhD": ["ph.d", "phd", "doctorate"],
    "M.Tech": ["m.tech", "mtech", "m. tech"], "M.E.": ["m.e.", "master of engineering"],
    "MBA": ["mba", "master of business"], "MCA": ["mca"], "M.Sc": ["m.sc", "msc", "master of science"],
    "B.Tech": ["b.tech", "btech", "b. tech"], "B.E.": ["b.e.", "b.e ", "bachelor of engineering"],
    "BCA": ["bca"], "B.Sc": ["b.sc", "bsc", "bachelor of science"], "BBA": ["bba"], "B.Com": ["b.com", "bcom"],
    "Diploma": ["diploma"],
}

SECTION_WORDS = {"resume", "curriculum vitae", "cv", "profile", "summary", "contact"}


class ResumeError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if name.endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        if name.endswith(".txt"):
            return data.decode("utf-8", errors="ignore")
    except Exception as exc:  # corrupted or encrypted files
        raise ResumeError(f"We couldn't read this file ({exc.__class__.__name__}). Upload a different copy.")
    raise ResumeError("Upload your resume as a PDF, DOCX or TXT file.")


def _guess_name(lines: list[str]) -> str:
    for line in lines[:8]:
        clean = line.strip()
        if not clean or EMAIL_RE.search(clean) or any(ch.isdigit() for ch in clean):
            continue
        if clean.lower() in SECTION_WORDS:
            continue
        words = clean.replace(",", " ").split()
        if 2 <= len(words) <= 4 and all(w.replace(".", "").isalpha() for w in words):
            return " ".join(w.capitalize() if w.isupper() else w for w in words)
    return ""


def _experience_years(text: str) -> float:
    stated = [float(m) for m in YEARS_RE.findall(text) if float(m) <= 45]
    if stated:
        return max(stated)
    this_year = date.today().year
    total = 0
    for start, end in RANGE_RE.findall(text):
        start_year = int(start)
        end_year = this_year if not end[:2].isdigit() else int(end)
        if 0 <= end_year - start_year <= 45:
            total += end_year - start_year
    return float(min(total, 45))


def parse_resume(text: str) -> dict:
    lines = [line for line in text.splitlines() if line.strip()]
    lowered = text.lower()
    email = EMAIL_RE.search(text)
    phone = next((p.strip() for p in PHONE_RE.findall(text) if len(re.sub(r"\D", "", p)) >= 10), "")
    education = [degree for degree, keys in DEGREES.items()
                 if any(re.search(r"(?<![a-z])" + re.escape(k.strip()) + r"(?![a-z])", lowered) for k in keys)]
    return {
        "name": _guess_name(lines),
        "email": email.group(0) if email else "",
        "phone": phone,
        "skills": extract_skills(text),
        "experience_years": _experience_years(text),
        "education": education,
        "links": sorted({m.lower() for m in LINK_RE.findall(text)}),
        "word_count": len(text.split()),
    }
